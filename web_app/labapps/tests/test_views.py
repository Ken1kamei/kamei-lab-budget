import hashlib
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
import zipfile

import pytest
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from numbers_parser import Document
from openpyxl import Workbook

from budget.models import (
    CategoryAllocation,
    FiscalYear,
    InvoiceDraft,
    LabMember,
    Team,
    Transaction,
)
from labapps.models import KnowledgeRecord, LabAppAudit, SheetRecord, SlackConnection
from labapps.services.knowledge_catalog import refresh_knowledge_indexes
from labapps.services.members import (
    delete_member_record,
    member_reference_summary,
    remove_member_access,
    sync_registry_member_mirror,
    upsert_member_access,
)
from labapps.tests.test_knowledge import protocol_docx_bytes


pytestmark = pytest.mark.django_db


def add_record(table, record_id, payload, source="registry"):
    SheetRecord.objects.create(
        source=source, table_name=table, record_id=record_id, payload=payload
    )


def fake_registry_upsert(table_name, payload, **kwargs):
    key = {
        "Members": "member_id",
        "Teams": "team_id",
        "App_Roles": "app_role_id",
        "Member_Teams": "member_team_id",
    }[table_name]
    SheetRecord.objects.update_or_create(
        source="registry",
        table_name=table_name,
        record_id=payload[key],
        defaults={"payload": payload},
    )
    return payload


def fake_registry_replace(table_name, rows, **kwargs):
    key = {
        "Members": "member_id",
        "Teams": "team_id",
        "App_Roles": "app_role_id",
        "Member_Teams": "member_team_id",
        "Projects": "project_id",
        "Milestones": "milestone_id",
        "Experiments": "experiment_id",
    }[table_name]
    source = (
        "tracker"
        if table_name in {"Projects", "Milestones", "Experiments"}
        else "registry"
    )
    SheetRecord.objects.filter(source=source, table_name=table_name).delete()
    for index, row in enumerate(rows, start=1):
        SheetRecord.objects.create(
            source=source,
            table_name=table_name,
            record_id=row.get(key) or f"legacy-{table_name}-{index}",
            payload=row,
        )
    return rows


def test_new_member_id_skips_dangling_references_and_deleted_member_audit(
    monkeypatch,
):
    seed_pi()
    add_record(
        "App_Roles",
        "AR011",
        {"app_role_id": "AR011", "member_id": "M011", "active": "FALSE"},
    )
    add_record(
        "Audit_Log",
        "AUD014",
        {
            "audit_id": "AUD014",
            "target_type": "Member",
            "target_id": "M014",
        },
    )
    monkeypatch.setattr(
        "labapps.services.members.upsert_record", fake_registry_upsert
    )
    monkeypatch.setattr(
        "labapps.services.members.append_registry_audit", lambda **kwargs: None
    )

    payload = upsert_member_access(
        {
            "member_id": "",
            "email": "new.member@nyu.edu",
            "name": "New Member",
            "display_name": "New Member",
            "global_role": "member",
            "active": True,
            "notes": "",
        },
        actor="kk4801@nyu.edu",
    )

    assert payload["member_id"] == "M015"
    assert SheetRecord.objects.get(
        table_name="Members", record_id="M015"
    ).payload["email"] == "new.member@nyu.edu"


def signed_in_client():
    user = get_user_model().objects.create_user(
        username="kk4801@nyu.edu", email="kk4801@nyu.edu"
    )
    client = Client()
    client.force_login(user)
    return client


def client_for(email):
    user = get_user_model().objects.create_user(username=email, email=email)
    client = Client()
    client.force_login(user)
    return client


def gantt_upload():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Gantt Import"
    sheet.append(
        [
            "Phase",
            "Task",
            "Assigned to",
            "Start Date",
            "End Date",
            "Progress %",
            "Status",
            "Next Action",
        ]
    )
    sheet.append(
        [
            "Planning",
            "Define scope",
            "kk4801@nyu.edu",
            "2026-09-01",
            "2026-09-05",
            50,
            "In progress",
            "Review scope",
        ]
    )
    buffer = BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(
        "project-gantt.xlsx",
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def invalid_gantt_upload():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Gantt Import"
    sheet.append(["Task", "Start Date", "End Date"])
    sheet.append(["Impossible task", "2026-09-10", "2026-09-01"])
    buffer = BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(
        "invalid-gantt.xlsx",
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def gantt_csv_upload():
    return SimpleUploadedFile(
        "project-gantt.csv",
        (
            "Phase,Task,Assigned to,Start Date,End Date,Progress %,Status,Next Action\n"
            "Planning,CSV scope,kk4801@nyu.edu,2026-09-01,2026-09-05,50,In progress,Review scope\n"
        ).encode(),
        content_type="text/csv",
    )


def gantt_numbers_upload(tmp_path):
    path = tmp_path / "project-gantt.numbers"
    document = Document(
        sheet_name="Gantt Import",
        table_name="Gantt Import",
        num_header_rows=1,
        num_header_cols=0,
        num_rows=2,
        num_cols=8,
    )
    table = document.sheets[0].tables[0]
    rows = [
        [
            "Phase",
            "Task",
            "Assigned to",
            "Start Date",
            "End Date",
            "Progress %",
            "Status",
            "Next Action",
        ],
        [
            "Planning",
            "Numbers scope",
            "kk4801@nyu.edu",
            "2026-09-01",
            "2026-09-05",
            50,
            "In progress",
            "Review scope",
        ],
    ]
    for row_number, row in enumerate(rows):
        for column_number, value in enumerate(row):
            table.write(row_number, column_number, value)
    document.save(path)
    return SimpleUploadedFile(
        path.name,
        path.read_bytes(),
        content_type="application/vnd.apple.numbers",
    )


def seed_pi():
    add_record(
        "Members", "M001",
        {"member_id": "M001", "email": "kk4801@nyu.edu", "display_name": "Ken", "global_role": "pi", "active": "TRUE"},
    )
    for app_id in ["budget", "project_tracker", "notebooks_protocols"]:
        add_record(
            "App_Roles", f"AR-{app_id}",
            {"member_id": "M001", "app_id": app_id, "app_role": "owner", "scope_team_id": "", "active": "TRUE"},
        )


def test_portal_tracker_and_knowledge_pages_render():
    seed_pi()
    add_record("Projects", "P001", {"project_id": "P001", "project": "Chip study", "owner_member_id": "M001"}, source="tracker")
    KnowledgeRecord.objects.create(
        record_id="P-0001", record_type="protocol", title="GSIS", team="Diabetes",
        metadata={"overview": ["Prepare buffer"]},
    )
    client = signed_in_client()
    portal = client.get("/portal/")
    tracker = client.get("/tracker/")
    knowledge = client.get("/knowledge/")
    assert portal.status_code == tracker.status_code == knowledge.status_code == 200
    assert b"Kamei Lab Apps" in portal.content
    assert b'class="sidebar"' not in portal.content
    assert b"Chip study" in tracker.content
    assert b'href="/tracker/#projects"' in tracker.content
    assert b">Transactions<" not in tracker.content
    assert b">Notebooks / protocols<" not in tracker.content
    assert b"Gantt chart" in tracker.content
    assert b"Kamei_Lab_Gantt_Import_Template.xlsx" in tracker.content
    assert b"Prepare buffer" not in knowledge.content
    assert b"Notebook registry" not in knowledge.content
    assert b"Find a record" in knowledge.content
    assert b'href="/knowledge/?type=protocol&amp;browse=1#library"' in knowledge.content
    assert b'href="/knowledge/?type=notebook&amp;browse=1#library"' in knowledge.content
    assert b'href="/knowledge/#search"' in knowledge.content
    assert b'href="/knowledge/upload/"' in knowledge.content
    assert b"Kamei_Lab_Protocol_Template.docx" in knowledge.content
    assert b"Download protocol template" in knowledge.content
    assert b">Transactions<" not in knowledge.content

    protocol = client.get("/knowledge/?record=P-0001")
    assert protocol.status_code == 200
    assert b"Prepare buffer" in protocol.content


def test_tracker_milestones_are_filtered_by_selected_project():
    seed_pi()
    for project_id, project_name in (
        ("P001", "First project"),
        ("P002", "Selected project"),
    ):
        add_record(
            "Projects",
            project_id,
            {
                "project_id": project_id,
                "project": project_name,
                "owner_member_id": "M001",
            },
            source="tracker",
        )
        add_record(
            "Milestones",
            f"MS-{project_id}",
            {
                "milestone_id": f"MS-{project_id}",
                "project_id": project_id,
                "project": project_name,
                "milestone": f"{project_name} milestone",
                "owner_member_id": "M001",
                "start_date": "2026-09-01",
                "due_date": "2026-09-05",
                "status": "In progress",
                "review_status": "Approved",
                "progress_percent": "40",
            },
            source="tracker",
        )

    response = signed_in_client().get("/tracker/?project=P002#milestones")

    assert response.status_code == 200
    assert response.context["selected_project"]["project_id"] == "P002"
    assert [
        row["milestone_id"] for row in response.context["selected_milestones"]
    ] == ["MS-P002"]
    assert b'data-project-id="P002"' in response.content
    assert b'data-milestone-id="MS-P002"' in response.content
    assert b'data-milestone-id="MS-P001"' not in response.content


def test_tracker_review_queue_shows_four_then_collapses_the_remainder():
    seed_pi()
    add_record(
        "Projects",
        "P001",
        {
            "project_id": "P001",
            "project": "Review project",
            "owner_member_id": "M001",
        },
        source="tracker",
    )
    for number in range(1, 7):
        milestone_id = f"MS-REVIEW-{number}"
        add_record(
            "Milestones",
            milestone_id,
            {
                "milestone_id": milestone_id,
                "project_id": "P001",
                "project": "Review project",
                "milestone": f"Review milestone {number}",
                "owner_member_id": "M001",
                "status": "In progress",
                "review_status": "Pending",
                "next_action": "Review this milestone",
            },
            source="tracker",
        )

    response = signed_in_client().get("/tracker/?project=P001#review")

    assert response.status_code == 200
    assert response.content.count(b'data-review-group="initial"') == 4
    assert response.content.count(b'data-review-group="overflow"') == 2
    assert b'data-review-overflow-count="2"' in response.content
    assert b"Show 2 more" in response.content
    assert b"Show fewer" in response.content


@patch("labapps.views.append_history")
@patch("labapps.views.upsert_record")
def test_milestone_progress_update_returns_to_project_and_updates_gantt(
    mock_upsert,
    mock_history,
):
    seed_pi()
    add_record(
        "Projects",
        "P002",
        {
            "project_id": "P002",
            "project": "Progress project",
            "owner_member_id": "M001",
        },
        source="tracker",
    )
    add_record(
        "Milestones",
        "MS-PROGRESS",
        {
            "milestone_id": "MS-PROGRESS",
            "project_id": "P002",
            "project": "Progress project",
            "milestone": "Measure progress",
            "owner_member_id": "M001",
            "start_date": "2026-09-01",
            "due_date": "2026-09-05",
            "status": "Not started",
            "review_status": "Approved",
            "next_action": "Begin",
            "progress_percent": "0",
        },
        source="tracker",
    )
    client = signed_in_client()

    update = client.post(
        "/tracker/?project=P002",
        {
            "action": "update",
            "table_name": "Milestones",
            "record_id": "MS-PROGRESS",
            "status": "In progress",
            "progress_percent": "65",
            "next_action": "Continue",
            "blocker_reason": "",
            "help_needed_from": "",
            "update_note": "Weekly update",
        },
    )

    assert update.status_code == 302
    assert update["Location"] == "/tracker/?project=P002#milestones"
    saved = mock_upsert.call_args.args[1]
    assert saved["project_id"] == "P002"
    assert saved["status"] == "In progress"
    assert saved["progress_percent"] == "65.0"
    assert mock_history.called

    record = SheetRecord.objects.get(
        source="tracker",
        table_name="Milestones",
        record_id="MS-PROGRESS",
    )
    record.payload = saved
    record.save(update_fields=["payload"])
    rendered = client.get("/tracker/?project=P002")

    assert rendered.status_code == 200
    assert rendered.context["gantt"]["rows"][0]["progress_width"] == 65.0
    assert b"--progress-width: 65.0%" in rendered.content
    assert b"65.0%" in rendered.content


def test_portal_uses_integrated_routes_even_with_legacy_registry_urls():
    seed_pi()
    legacy_urls = {
        "budget": "https://legacy-budget.example.streamlit.app/",
        "project_tracker": "https://legacy-tracker.example.streamlit.app/",
        "notebooks_protocols": "https://legacy-knowledge.example.streamlit.app/",
    }
    for index, (app_id, app_url) in enumerate(legacy_urls.items(), start=1):
        add_record(
            "Apps",
            f"APP-{index}",
            {
                "app_id": app_id,
                "app_name": app_id,
                "app_url": app_url,
                "description": f"{app_id} description",
                "active": "TRUE",
            },
        )

    response = signed_in_client().get("/portal/")

    assert response.status_code == 200
    assert b"streamlit.app" not in response.content
    assert b'href="/"' in response.content
    assert b'href="/tracker/"' in response.content
    assert b'href="/knowledge/"' in response.content


def test_portal_action_panel_aggregates_tracker_and_budget_attention():
    seed_pi()
    add_record(
        "Milestones",
        "MS-OVERDUE",
        {
            "milestone_id": "MS-OVERDUE",
            "milestone": "Overdue work",
            "due_date": "2000-01-01",
            "status": "Blocked",
            "review_status": "Pending",
        },
        source="tracker",
    )
    add_record(
        "Experiments",
        "EXP-DONE",
        {
            "experiment_id": "EXP-DONE",
            "experiment_title": "Finished work",
            "due_date": "2000-01-01",
            "status": "Completed",
            "review_status": "Approved",
        },
        source="tracker",
    )
    add_record(
        "Experiments",
        "EXP-BAD-DATE",
        {
            "experiment_id": "EXP-BAD-DATE",
            "experiment_title": "No valid deadline",
            "due_date": "not-a-date",
            "status": "Running",
            "review_status": "Approved",
        },
        source="tracker",
    )
    fiscal_year = FiscalYear.objects.create(label="FY2026-27")
    LabMember.objects.create(
        email="kk4801@nyu.edu", display_name="Ken", highest_role="pi"
    )
    Team.objects.create(
        fiscal_year=fiscal_year,
        name="Diabetes",
        allocation_usd=Decimal("1000"),
    )
    CategoryAllocation.objects.create(
        fiscal_year=fiscal_year,
        category="Consumables",
        budget_usd=Decimal("1000"),
    )
    Transaction.objects.create(
        fiscal_year=fiscal_year,
        transaction_id="TX-1",
        team="Diabetes",
        category="Consumables",
        amount_usd_equiv=Decimal("900"),
    )
    InvoiceDraft.objects.create(
        uploader_email="member@nyu.edu",
        file_name="invoice.pdf",
        file_sha256="a" * 64,
        fiscal_year=fiscal_year,
        status="review",
    )

    response = signed_in_client().get("/portal/")

    assert response.status_code == 200
    content = response.content.decode()
    assert "Needs attention" in content
    assert "Overdue" in content
    assert "Blocked" in content
    assert content.index("Workspace") < content.index("Across your apps")
    assert content.index("Across your apps") < content.index("Active members")
    assert content.index("Active members") < content.index("This week")
    assert "Tracker 1 · invoices 1" in content
    assert "2 critical · highest 90%" not in content
    assert "0 critical · highest 90%" in content
    action_counts = {
        item["label"]: item["count"] for item in response.context["action_panel"]["items"]
    }
    assert action_counts == {
        "Overdue": 1,
        "Blocked": 1,
        "Pending approval": 2,
        "Budget alert": 2,
    }


def test_protocol_template_is_valid_and_linked_from_upload_page():
    seed_pi()
    template = (
        Path(__file__).resolve().parents[2]
        / "labapps/static/labapps/Kamei_Lab_Protocol_Template.docx"
    )
    assert template.exists()
    assert zipfile.is_zipfile(template)

    response = signed_in_client().get("/knowledge/upload/")
    assert response.status_code == 200
    assert b"Kamei_Lab_Protocol_Template.docx" in response.content
    assert b"Download protocol template" in response.content


def test_protocol_template_is_available_to_read_only_members():
    add_record(
        "Members",
        "M002",
        {
            "member_id": "M002",
            "email": "member@nyu.edu",
            "display_name": "Lab member",
            "global_role": "member",
            "active": "TRUE",
        },
    )
    add_record(
        "App_Roles",
        "AR-notebooks-reader",
        {
            "member_id": "M002",
            "app_id": "notebooks_protocols",
            "app_role": "viewer",
            "scope_team_id": "",
            "active": "TRUE",
        },
    )

    response = client_for("member@nyu.edu").get("/knowledge/")
    assert response.status_code == 200
    assert b"Kamei_Lab_Protocol_Template.docx" in response.content
    assert b"Download protocol template" in response.content
    assert b'href="/knowledge/upload/"' not in response.content


@patch("labapps.views.replace_project_gantt")
def test_gantt_upload_previews_then_replaces_only_imported_project_rows(mock_replace):
    seed_pi()
    add_record(
        "Projects",
        "P001",
        {
            "project_id": "P001",
            "project": "Chip study",
            "aim": "Model disease",
            "owner_member_id": "M001",
        },
        source="tracker",
    )
    add_record(
        "Milestones",
        "MS001",
        {
            "milestone_id": "MS001",
            "project_id": "P001",
            "milestone": "Manual milestone",
            "owner_member_id": "M001",
        },
        source="tracker",
    )
    add_record(
        "Milestones",
        "MS-GANTT-OLD",
        {
            "milestone_id": "MS-GANTT-OLD",
            "project_id": "P001",
            "milestone": "Previous import",
            "owner_member_id": "M001",
        },
        source="tracker",
    )
    add_record(
        "Milestones",
        "MS-GANTT-OTHER",
        {
            "milestone_id": "MS-GANTT-OTHER",
            "project_id": "P002",
            "milestone": "Another project",
            "owner_member_id": "M001",
        },
        source="tracker",
    )
    client = signed_in_client()

    preview = client.post(
        "/tracker/",
        {
            "action": "gantt_preview",
            "gantt-project_id": "P001",
            "gantt-default_owner_member_id": "M001",
            "gantt-gantt_file": gantt_upload(),
        },
    )

    assert preview.status_code == 200
    assert b"Gantt chart preview" in preview.content
    assert b"Define scope" in preview.content
    assert b"Save 1 task and show chart" in preview.content
    assert b'class="gantt-track"' in preview.content
    assert b"Review imported task details (1)" in preview.content
    assert preview.content.index(b"Save 1 task and show chart") < preview.content.index(
        b'class="gantt-preview-chart"'
    ) < preview.content.index(b"Review imported task details (1)")
    stored = client.session["gantt_import_preview"]
    assert stored["project_id"] == "P001"
    assert stored["actor"] == "kk4801@nyu.edu"
    assert len(stored["rows"]) == 1

    confirm = client.post(
        "/tracker/",
        {
            "action": "gantt_confirm",
            "preview_token": stored["token"],
        },
    )

    assert confirm.status_code == 302
    assert "gantt_project=P001" in confirm["Location"]
    assert mock_replace.call_args.args[0] == "P001"
    saved_rows = mock_replace.call_args.args[1]
    assert any(
        row["milestone"] == "Define scope"
        and row["milestone_id"].startswith("MS-GANTT-")
        for row in saved_rows
    )


@pytest.mark.parametrize(
    ("upload_factory", "expected_task"),
    [
        (lambda tmp_path: gantt_csv_upload(), "CSV scope"),
        (gantt_numbers_upload, "Numbers scope"),
    ],
)
@patch("labapps.views.replace_project_gantt")
def test_csv_and_numbers_gantt_uploads_use_the_preview_flow(
    mock_replace,
    tmp_path,
    upload_factory,
    expected_task,
):
    seed_pi()
    add_record(
        "Projects",
        "P001",
        {
            "project_id": "P001",
            "project": "Chip study",
            "aim": "Model disease",
            "owner_member_id": "M001",
        },
        source="tracker",
    )
    client = signed_in_client()

    preview = client.post(
        "/tracker/",
        {
            "action": "gantt_preview",
            "gantt-project_id": "P001",
            "gantt-default_owner_member_id": "M001",
            "gantt-gantt_file": upload_factory(tmp_path),
        },
    )

    assert preview.status_code == 200
    assert b"Gantt chart preview" in preview.content
    assert expected_task.encode() in preview.content
    assert b"Save 1 task and show chart" in preview.content
    assert b'class="gantt-track"' in preview.content
    stored = client.session["gantt_import_preview"]
    assert stored["project_id"] == "P001"
    assert stored["rows"][0]["milestone"] == expected_task

    confirmed = client.post(
        "/tracker/",
        {
            "action": "gantt_confirm",
            "preview_token": stored["token"],
        },
    )

    assert confirmed.status_code == 302
    mock_replace.assert_called_once()
    assert mock_replace.call_args.args[1][0]["milestone"] == expected_task


@patch("labapps.views.replace_project_gantt")
def test_invalid_gantt_preview_cannot_be_confirmed(mock_replace):
    seed_pi()
    add_record(
        "Projects",
        "P001",
        {
            "project_id": "P001",
            "project": "Chip study",
            "owner_member_id": "M001",
        },
        source="tracker",
    )
    client = signed_in_client()

    preview = client.post(
        "/tracker/",
        {
            "action": "gantt_preview",
            "gantt-project_id": "P001",
            "gantt-default_owner_member_id": "M001",
            "gantt-gantt_file": invalid_gantt_upload(),
        },
    )

    assert preview.status_code == 200
    assert b"ends before its start date" in preview.content
    assert "gantt_import_preview" not in client.session
    assert b"Confirm and save to Google Sheets" not in preview.content
    mock_replace.assert_not_called()


@patch("labapps.views.replace_project_gantt")
def test_read_only_project_tracker_role_cannot_upload_gantt(mock_replace):
    add_record(
        "Members",
        "M004",
        {
            "member_id": "M004",
            "email": "viewer@nyu.edu",
            "display_name": "Viewer",
            "active": "TRUE",
        },
    )
    add_record(
        "App_Roles",
        "AR-viewer",
        {
            "member_id": "M004",
            "app_id": "project_tracker",
            "app_role": "viewer",
            "scope_team_id": "",
            "active": "TRUE",
        },
    )
    add_record(
        "Projects",
        "P001",
        {
            "project_id": "P001",
            "project": "Chip study",
            "owner_member_id": "M004",
        },
        source="tracker",
    )
    client = client_for("viewer@nyu.edu")

    response = client.post(
        "/tracker/",
        {
            "action": "gantt_preview",
            "gantt-project_id": "P001",
            "gantt-default_owner_member_id": "M004",
            "gantt-gantt_file": gantt_upload(),
        },
    )

    assert response.status_code == 403
    assert "gantt_import_preview" not in client.session
    mock_replace.assert_not_called()


def test_knowledge_keyword_search_matches_notebooks_and_protocol_content():
    seed_pi()
    KnowledgeRecord.objects.create(
        record_id="P-0001", record_type="protocol", title="GSIS workflow",
        team="Diabetes", metadata={"overview": ["Prepare assay buffer"]},
    )
    KnowledgeRecord.objects.create(
        record_id="N-0001", record_type="notebook", title="Buffer optimization",
        owner="Satoshi", team="IoC", original_filename="buffer-notes.pdf",
    )
    KnowledgeRecord.objects.create(
        record_id="N-0002", record_type="notebook", title="Unrelated imaging log",
        owner="Maab", team="IoC",
    )

    client = signed_in_client()
    response = client.get("/knowledge/?q=buffer")

    assert response.status_code == 200
    assert response.context["search_total"] == 2
    assert {
        row.record_id for row in response.context["search_results"]
    } == {"P-0001", "N-0001"}
    assert b"GSIS workflow" in response.content
    assert b"Buffer optimization" in response.content
    assert b"Notebook registry" not in response.content


def test_knowledge_available_view_restores_legacy_statuses_without_auto_listing():
    seed_pi()
    KnowledgeRecord.objects.create(
        record_id="P-CANDIDATE",
        record_type="protocol",
        title="Legacy candidate protocol",
        status="candidate",
    )
    KnowledgeRecord.objects.create(
        record_id="N-INDEXED",
        record_type="notebook",
        title="Legacy indexed notebook",
        status="indexed",
    )
    client = signed_in_client()

    overview = client.get("/knowledge/")
    protocol_list = client.get("/knowledge/?type=protocol&browse=1")
    notebook_list = client.get("/knowledge/?type=notebook&browse=1")

    assert overview.status_code == 200
    assert overview.context["counts"] == {
        "protocols": 1,
        "notebooks": 1,
        "duplicates": 0,
    }
    assert overview.context["show_results"] is False
    assert list(overview.context["search_results"]) == []
    assert [
        row.record_id for row in protocol_list.context["search_results"]
    ] == ["P-CANDIDATE"]
    assert [
        row.record_id for row in notebook_list.context["search_results"]
    ] == ["N-INDEXED"]


def test_verified_duplicate_records_are_grouped_and_alias_opens_canonical():
    seed_pi()
    canonical = KnowledgeRecord.objects.create(
        record_id="N-0001",
        record_type="notebook",
        title="MEF notebook",
        status="indexed",
        content_sha256="a" * 64,
        metadata={"summary": ["Canonical notebook content"]},
    )
    alias = KnowledgeRecord.objects.create(
        record_id="N-0002",
        record_type="notebook",
        title="MEF notebook copy",
        status="indexed",
        content_sha256="a" * 64,
        source_path="/source/copy.pdf",
    )
    refresh_knowledge_indexes()
    canonical.refresh_from_db()
    alias.refresh_from_db()
    client = signed_in_client()

    browse = client.get("/knowledge/?type=notebook&browse=1")
    detail = client.get(f"/knowledge/?record={alias.record_id}")

    assert canonical.canonical_record_id == canonical.record_id
    assert alias.canonical_record_id == canonical.record_id
    assert browse.context["counts"]["duplicates"] == 1
    assert [
        row.record_id for row in browse.context["search_results"]
    ] == [canonical.record_id]
    assert detail.context["selected_record"].record_id == canonical.record_id
    assert [
        row.record_id for row in detail.context["selected_aliases"]
    ] == [alias.record_id]
    assert b"Canonical notebook content" in detail.content
    assert b"verified duplicate source" in detail.content


def test_notebook_structured_content_uses_common_viewer():
    seed_pi()
    record = KnowledgeRecord.objects.create(
        record_id="N-LABBOOK",
        record_type="notebook",
        title="Differentiation notebook",
        status="indexed",
        original_filename="notebook.pdf",
        metadata={
            "sections": [
                {
                    "heading": "Day 3 observations",
                    "blocks": [
                        {
                            "kind": "paragraph",
                            "text": "Activin A treatment completed.",
                        }
                    ],
                }
            ]
        },
    )

    response = signed_in_client().get(
        f"/knowledge/?record={record.record_id}"
    )

    assert response.status_code == 200
    assert response.context["selected_record"].record_type == "notebook"
    assert b"Day 3 observations" in response.content
    assert b"Activin A treatment completed." in response.content


@patch("labapps.views.open_knowledge_file", return_value=BytesIO(b"%PDF-demo"))
def test_previewable_notebook_original_opens_inline(mock_open):
    seed_pi()
    record = KnowledgeRecord.objects.create(
        record_id="N-PDF",
        record_type="notebook",
        title="PDF notebook",
        status="active",
        object_name="knowledge/N-PDF/notebook.pdf",
        original_filename="notebook.pdf",
        metadata={"content_type": "application/pdf"},
    )

    response = signed_in_client().get(
        f"/knowledge/{record.record_id}/original/"
    )

    assert response.status_code == 200
    assert response["Content-Type"] == "application/pdf"
    assert response["Content-Disposition"].startswith("inline;")
    assert response["X-Content-Type-Options"] == "nosniff"
    assert b"".join(response.streaming_content) == b"%PDF-demo"
    mock_open.assert_called_once_with(record.object_name)


def test_knowledge_results_are_paginated_twenty_at_a_time():
    seed_pi()
    KnowledgeRecord.objects.bulk_create(
        [
            KnowledgeRecord(
                record_id=f"N-{index:04d}",
                record_type="notebook",
                title=f"Notebook {index:04d}",
                status="indexed",
            )
            for index in range(25)
        ]
    )

    response = signed_in_client().get(
        "/knowledge/?type=notebook&browse=1"
    )

    assert response.context["search_total"] == 25
    assert len(response.context["search_results"]) == 20
    assert response.context["search_page"].paginator.num_pages == 2


def test_knowledge_record_can_be_archived_and_restored():
    seed_pi()
    record = KnowledgeRecord.objects.create(
        record_id="P-LIFECYCLE",
        record_type="protocol",
        title="Lifecycle protocol",
        status="active",
    )
    client = signed_in_client()

    archived = client.post(
        f"/knowledge/{record.record_id}/status/",
        {"status": "archived"},
    )

    assert archived.status_code == 302
    record.refresh_from_db()
    assert record.status == "archived"
    active_page = client.get("/knowledge/")
    archived_page = client.get("/knowledge/?status=archived&browse=1")
    assert list(active_page.context["protocols"]) == []
    assert [
        row.record_id for row in archived_page.context["search_results"]
    ] == [
        "P-LIFECYCLE"
    ]
    audit = LabAppAudit.objects.get(
        target=record.record_id,
        action="status_updated",
    )
    assert audit.before == {"status": "active"}
    assert audit.after == {"status": "archived"}


def test_portal_member_save_updates_budget_iap_allowlist(
    monkeypatch,
    settings,
):
    settings.ENABLE_SHEET_WRITES = True
    settings.SHEET_WRITE_ALLOWED_EMAILS = {"*"}
    seed_pi()

    monkeypatch.setattr(
        "labapps.services.members.upsert_record", fake_registry_upsert
    )
    monkeypatch.setattr(
        "labapps.services.members.append_registry_audit", lambda **kwargs: None
    )
    client = signed_in_client()

    response = client.post(
        "/portal/admin/",
        {
            "action": "member",
            "member-email": "new.member@nyu.edu",
            "member-name": "New Member",
            "member-display_name": "New Member",
            "member-global_role": "member",
            "member-active": "on",
            "member-notes": "",
        },
    )

    assert response.status_code == 302
    allowlisted = LabMember.objects.get(email="new.member@nyu.edu")
    assert allowlisted.display_name == "New Member"
    assert allowlisted.highest_role == "member"
    assert allowlisted.active is True


def test_portal_member_demotion_revokes_budget_pi_role(
    monkeypatch,
    settings,
):
    settings.ENABLE_SHEET_WRITES = True
    settings.SHEET_WRITE_ALLOWED_EMAILS = {"*"}
    seed_pi()
    LabMember.objects.create(
        email="former.pi@nyu.edu",
        display_name="Former PI",
        highest_role="pi",
        active=True,
    )
    add_record(
        "Members",
        "M009",
        {
            "member_id": "M009",
            "email": "former.pi@nyu.edu",
            "name": "Former PI",
            "display_name": "Former PI",
            "global_role": "pi",
            "active": "TRUE",
        },
    )

    monkeypatch.setattr(
        "labapps.services.members.upsert_record", fake_registry_upsert
    )
    monkeypatch.setattr(
        "labapps.services.members.append_registry_audit", lambda **kwargs: None
    )

    response = signed_in_client().post(
        "/portal/admin/",
        {
            "action": "member",
            "member-email": "former.pi@nyu.edu",
            "member-name": "Former PI",
            "member-display_name": "Former PI",
            "member-global_role": "member",
            "member-active": "on",
            "member-notes": "",
        },
    )

    assert response.status_code == 302
    assert LabMember.objects.get(
        email="former.pi@nyu.edu"
    ).highest_role == "member"


def test_portal_active_members_tile_links_to_member_management_for_pi():
    seed_pi()

    response = signed_in_client().get("/portal/")

    assert response.status_code == 200
    assert b'href="/portal/admin/#members"' in response.content
    assert b"Manage members" in response.content


def test_portal_member_form_accepts_only_nyu_or_approved_lab_account(
    monkeypatch,
    settings,
):
    settings.ENABLE_SHEET_WRITES = True
    settings.SHEET_WRITE_ALLOWED_EMAILS = {"*"}
    settings.LAB_MEMBER_EMAIL_EXCEPTIONS = {"nyuadkameilab@gmail.com"}
    seed_pi()
    monkeypatch.setattr(
        "labapps.services.members.upsert_record", fake_registry_upsert
    )
    monkeypatch.setattr(
        "labapps.services.members.append_registry_audit", lambda **kwargs: None
    )
    client = signed_in_client()
    base = {
        "action": "member",
        "member-name": "Kamei Lab",
        "member-display_name": "Kamei Lab",
        "member-global_role": "member",
        "member-active": "on",
        "member-notes": "",
    }

    rejected = client.post(
        "/portal/admin/",
        {**base, "member-email": "unapproved@gmail.com"},
    )
    accepted = client.post(
        "/portal/admin/",
        {**base, "member-email": "NYUADKAMEILAB@gmail.com"},
    )

    assert rejected.status_code == 200
    assert b"Use an @nyu.edu account" in rejected.content
    assert accepted.status_code == 302
    assert LabMember.objects.get(email="nyuadkameilab@gmail.com").active is True


def test_registry_member_edit_updates_same_record_and_web_mirror(
    monkeypatch, settings
):
    settings.ENABLE_SHEET_WRITES = True
    settings.SHEET_WRITE_ALLOWED_EMAILS = {"*"}
    seed_pi()
    add_record(
        "Members",
        "M002",
        {
            "member_id": "M002",
            "email": "wrong@nyu.edu",
            "name": "Wrong Name",
            "display_name": "Wrong",
            "global_role": "member",
            "active": "TRUE",
            "notes": "Needs correction",
        },
    )
    LabMember.objects.create(
        email="wrong@nyu.edu", display_name="Wrong", active=True
    )
    monkeypatch.setattr(
        "labapps.services.members.upsert_record", fake_registry_upsert
    )
    monkeypatch.setattr(
        "labapps.services.members.append_registry_audit", lambda **kwargs: None
    )
    client = signed_in_client()

    edit_page = client.get("/portal/admin/?edit_member=M002#members")
    response = client.post(
        "/portal/admin/",
        {
            "action": "member",
            "member-member_id": "M002",
            "member-email": "correct@nyu.edu",
            "member-name": "Correct Name",
            "member-display_name": "Correct",
            "member-global_role": "lead",
            "member-active": "on",
            "member-notes": "Corrected in Lab Registry",
        },
    )

    assert edit_page.status_code == 200
    assert b"Save member changes" in edit_page.content
    assert b'value="wrong@nyu.edu"' in edit_page.content
    assert response.status_code == 302
    member = SheetRecord.objects.get(
        source="registry", table_name="Members", record_id="M002"
    ).payload
    assert member["email"] == "correct@nyu.edu"
    assert member["display_name"] == "Correct"
    assert member["global_role"] == "lead"
    assert not LabMember.objects.filter(email="wrong@nyu.edu").exists()
    mirror = LabMember.objects.get(email="correct@nyu.edu")
    assert mirror.display_name == "Correct"
    assert mirror.highest_role == "lead"


def test_registry_member_edit_rejects_duplicate_email(monkeypatch, settings):
    settings.ENABLE_SHEET_WRITES = True
    settings.SHEET_WRITE_ALLOWED_EMAILS = {"*"}
    seed_pi()
    add_record(
        "Members",
        "M002",
        {
            "member_id": "M002",
            "email": "first@nyu.edu",
            "name": "First",
            "display_name": "First",
            "global_role": "member",
            "active": "TRUE",
        },
    )
    add_record(
        "Members",
        "M003",
        {
            "member_id": "M003",
            "email": "second@nyu.edu",
            "name": "Second",
            "display_name": "Second",
            "global_role": "member",
            "active": "TRUE",
        },
    )
    monkeypatch.setattr(
        "labapps.services.members.upsert_record", fake_registry_upsert
    )

    response = signed_in_client().post(
        "/portal/admin/",
        {
            "action": "member",
            "member-member_id": "M002",
            "member-email": "second@nyu.edu",
            "member-name": "First",
            "member-display_name": "First",
            "member-global_role": "member",
            "member-active": "on",
            "member-notes": "",
        },
    )

    assert response.status_code == 200
    assert b"already assigned to another member" in response.content
    assert SheetRecord.objects.get(
        source="registry", table_name="Members", record_id="M002"
    ).payload["email"] == "first@nyu.edu"


def test_registry_mirror_uses_central_names_roles_and_team_access():
    fiscal_year = FiscalYear.objects.create(label="FY2026-27")
    Team.objects.create(fiscal_year=fiscal_year, name="Core Lab", active=True)
    Team.objects.create(fiscal_year=fiscal_year, name="Diabetes", active=True)
    add_record(
        "Members",
        "M007",
        {
            "member_id": "M007",
            "email": "maab@nyu.edu",
            "name": "Maab Correct",
            "display_name": "Maab",
            "global_role": "admin",
            "active": "TRUE",
        },
    )
    add_record(
        "Teams",
        "T001",
        {"team_id": "T001", "team_name": "Core Lab", "active": "TRUE"},
    )
    add_record(
        "Teams",
        "T003",
        {"team_id": "T003", "team_name": "Diabetes", "active": "TRUE"},
    )
    add_record(
        "Member_Teams",
        "MT001",
        {
            "member_team_id": "MT001",
            "member_id": "M007",
            "team_id": "T001",
            "team_role": "member",
            "active": "TRUE",
        },
    )
    add_record(
        "Member_Teams",
        "MT002",
        {
            "member_team_id": "MT002",
            "member_id": "M007",
            "team_id": "T003",
            "team_role": "lead",
            "active": "TRUE",
        },
    )

    sync_registry_member_mirror("M007")

    member = LabMember.objects.get(email="maab@nyu.edu")
    assert member.display_name == "Maab"
    assert member.highest_role == "budget_manager"
    assert member.team_names == ["Core Lab", "Diabetes"]
    assert member.team_roles == {
        "FY2026-27": {"Core Lab": "member", "Diabetes": "lead"}
    }


def test_registry_shows_inactive_members_and_permanently_deletes_one(
    monkeypatch, settings
):
    settings.ENABLE_SHEET_WRITES = True
    settings.SHEET_WRITE_ALLOWED_EMAILS = {"*"}
    seed_pi()
    add_record(
        "Members",
        "M005",
        {
            "member_id": "M005",
            "email": "test@test.com",
            "name": "Test Member",
            "display_name": "Test Member",
            "global_role": "member",
            "active": "FALSE",
        },
    )
    add_record(
        "Member_Teams",
        "MT005",
        {
            "member_team_id": "MT005",
            "member_id": "M005",
            "team_id": "T001",
            "active": "FALSE",
        },
    )
    add_record(
        "App_Roles",
        "AR005",
        {
            "app_role_id": "AR005",
            "member_id": "M005",
            "app_id": "budget",
            "app_role": "viewer",
            "active": "FALSE",
        },
    )
    LabMember.objects.create(email="test@test.com", active=False)
    monkeypatch.setattr(
        "labapps.services.members.replace_table", fake_registry_replace
    )
    monkeypatch.setattr(
        "labapps.services.members.append_registry_audit", lambda **kwargs: None
    )
    client = signed_in_client()

    page = client.get("/portal/admin/#members")
    confirm_page = client.get("/portal/admin/?delete_member=M005#members")
    response = client.post(
        "/portal/admin/",
        {
            "action": "member_delete",
            "delete-member_id": "M005",
            "delete-confirm_email": "test@test.com",
        },
    )

    assert page.status_code == 200
    assert b"Test Member" in page.content
    assert b"Inactive" in page.content
    assert b"1 active" in page.content
    assert b"Delete" in page.content
    assert b"Permanent deletion" in confirm_page.content
    assert response.status_code == 302
    assert not SheetRecord.objects.filter(
        source="registry", table_name="Members", record_id="M005"
    ).exists()
    assert not SheetRecord.objects.filter(
        source="registry", table_name="Member_Teams", record_id="MT005"
    ).exists()
    assert not SheetRecord.objects.filter(
        source="registry", table_name="App_Roles", record_id="AR005"
    ).exists()
    assert not LabMember.objects.filter(email="test@test.com").exists()


def test_registry_permanent_delete_transfers_linked_records(
    monkeypatch, settings
):
    settings.ENABLE_SHEET_WRITES = True
    settings.SHEET_WRITE_ALLOWED_EMAILS = {"*"}
    add_record(
        "Members",
        "M001",
        {
            "member_id": "M001",
            "email": "owner@nyu.edu",
            "display_name": "Owner",
            "global_role": "admin",
            "active": "TRUE",
        },
    )
    add_record(
        "Members",
        "M020",
        {
            "member_id": "M020",
            "email": "linked@nyu.edu",
            "display_name": "Linked",
            "global_role": "member",
            "active": "FALSE",
        },
    )
    add_record(
        "Projects",
        "P020",
        {"project_id": "P020", "owner_member_id": "M020"},
        source="tracker",
    )
    add_record(
        "Milestones",
        "MS020",
        {"milestone_id": "MS020", "owner_member_id": "M020"},
        source="tracker",
    )
    add_record(
        "Experiments",
        "EXP020",
        {"experiment_id": "EXP020", "member_id": "M020"},
        source="tracker",
    )
    monkeypatch.setattr(
        "labapps.services.members.replace_table", fake_registry_replace
    )
    monkeypatch.setattr(
        "labapps.services.members.append_registry_audit", lambda **kwargs: None
    )

    with pytest.raises(ValueError, match="Choose a different active member"):
        delete_member_record(
            "M020", actor="admin@nyu.edu", confirm_email="linked@nyu.edu"
        )

    assert SheetRecord.objects.filter(
        source="registry", table_name="Members", record_id="M020"
    ).exists()

    delete_member_record(
        "M020",
        actor="admin@nyu.edu",
        confirm_email="linked@nyu.edu",
        reassign_to_member_id="M001",
    )

    assert not SheetRecord.objects.filter(
        source="registry", table_name="Members", record_id="M020"
    ).exists()
    assert SheetRecord.objects.get(
        source="tracker", table_name="Projects", record_id="P020"
    ).payload["owner_member_id"] == "M001"
    assert SheetRecord.objects.get(
        source="tracker", table_name="Milestones", record_id="MS020"
    ).payload["owner_member_id"] == "M001"
    assert SheetRecord.objects.get(
        source="tracker", table_name="Experiments", record_id="EXP020"
    ).payload["member_id"] == "M001"


def test_portal_member_remove_revokes_all_access_and_slack(
    monkeypatch,
    settings,
):
    settings.ENABLE_SHEET_WRITES = True
    settings.SHEET_WRITE_ALLOWED_EMAILS = {"*"}
    seed_pi()
    add_record(
        "Members",
        "M002",
        {
            "member_id": "M002",
            "email": "member@nyu.edu",
            "name": "Member",
            "display_name": "Member",
            "global_role": "member",
            "active": "TRUE",
        },
    )
    add_record(
        "App_Roles",
        "AR002",
        {
            "app_role_id": "AR002",
            "member_id": "M002",
            "app_id": "notebooks_protocols",
            "app_role": "member",
            "scope_team_id": "T001",
            "active": "TRUE",
        },
    )
    add_record(
        "Member_Teams",
        "MT002",
        {
            "member_team_id": "MT002",
            "member_id": "M002",
            "team_id": "T001",
            "team_role": "member",
            "active": "TRUE",
        },
    )
    LabMember.objects.create(
        email="member@nyu.edu",
        display_name="Member",
        highest_role="member",
        active=True,
    )
    SlackConnection.objects.create(
        portal_email="member@nyu.edu",
        slack_team_id="T1",
        slack_team_name="KameiLab_NYUAD",
        slack_user_id="U1",
        slack_user_name="Member",
        access_token_ciphertext="encrypted",
    )
    monkeypatch.setattr(
        "labapps.services.members.upsert_record", fake_registry_upsert
    )
    monkeypatch.setattr(
        "labapps.services.members.append_registry_audit", lambda **kwargs: None
    )

    response = signed_in_client().post(
        "/portal/admin/",
        {"action": "member_remove", "member_id": "M002"},
    )

    assert response.status_code == 302
    assert SheetRecord.objects.get(
        table_name="Members", record_id="M002"
    ).payload["active"] == "FALSE"
    assert SheetRecord.objects.get(
        table_name="App_Roles", record_id="AR002"
    ).payload["active"] == "FALSE"
    assert SheetRecord.objects.get(
        table_name="Member_Teams", record_id="MT002"
    ).payload["active"] == "FALSE"
    assert LabMember.objects.get(email="member@nyu.edu").active is False
    assert not SlackConnection.objects.filter(portal_email="member@nyu.edu").exists()


def test_portal_member_remove_allows_revoking_legacy_unapproved_email(
    monkeypatch,
    settings,
):
    settings.ENABLE_SHEET_WRITES = True
    settings.SHEET_WRITE_ALLOWED_EMAILS = {"*"}
    seed_pi()
    add_record(
        "Members",
        "M002",
        {
            "member_id": "M002",
            "email": "member@example.edu",
            "name": "Legacy member",
            "display_name": "Legacy member",
            "global_role": "member",
            "active": "TRUE",
        },
    )
    LabMember.objects.create(
        email="member@example.edu",
        display_name="Legacy member",
        highest_role="member",
        active=True,
    )
    monkeypatch.setattr(
        "labapps.services.members.upsert_record", fake_registry_upsert
    )
    monkeypatch.setattr(
        "labapps.services.members.append_registry_audit", lambda **kwargs: None
    )

    response = signed_in_client().post(
        "/portal/admin/",
        {"action": "member_remove", "member_id": "M002"},
        follow=True,
    )

    assert response.status_code == 200
    assert b"Legacy member" in response.content
    assert b"Inactive" in response.content
    assert SheetRecord.objects.get(
        table_name="Members", record_id="M002"
    ).payload["active"] == "FALSE"
    assert LabMember.objects.get(email="member@example.edu").active is False


def test_portal_member_remove_protects_pi_and_non_admin_access(settings):
    settings.PI_EMAIL = "kk4801@nyu.edu"
    seed_pi()

    with pytest.raises(ValueError, match="PI account cannot be removed"):
        remove_member_access("M001", actor="admin@nyu.edu")

    add_record(
        "Members",
        "M002",
        {
            "member_id": "M002",
            "email": "member@nyu.edu",
            "display_name": "Member",
            "global_role": "member",
            "active": "TRUE",
        },
    )
    response = client_for("member@nyu.edu").post(
        "/portal/admin/",
        {"action": "member_remove", "member_id": "M001"},
    )

    assert response.status_code == 403


def test_scoped_tracker_role_cannot_switch_to_another_team():
    add_record(
        "Members", "M002",
        {"member_id": "M002", "email": "lead@nyu.edu", "display_name": "Lead", "active": "TRUE"},
    )
    add_record(
        "Members", "M003",
        {"member_id": "M003", "email": "other@nyu.edu", "display_name": "Other", "active": "TRUE"},
    )
    add_record("Teams", "T001", {"team_id": "T001", "team_name": "IoC", "active": "TRUE"})
    add_record("Teams", "T002", {"team_id": "T002", "team_name": "Diabetes", "active": "TRUE"})
    add_record(
        "Member_Teams", "MT002",
        {"member_team_id": "MT002", "member_id": "M002", "team_id": "T002", "active": "TRUE"},
    )
    add_record(
        "Member_Teams", "MT003",
        {"member_team_id": "MT003", "member_id": "M003", "team_id": "T001", "active": "TRUE"},
    )
    add_record(
        "App_Roles", "AR002",
        {"member_id": "M002", "app_id": "project_tracker", "app_role": "lead", "scope_team_id": "T002", "active": "TRUE"},
    )
    add_record(
        "Projects", "P001",
        {"project_id": "P001", "project": "Other team project", "owner_member_id": "M003"},
        source="tracker",
    )
    add_record(
        "Projects", "P002",
        {"project_id": "P002", "project": "Scoped project", "owner_member_id": "M002"},
        source="tracker",
    )
    add_record(
        "Milestones", "MS001",
        {
            "milestone_id": "MS001", "project_id": "P001",
            "milestone": "Other team milestone", "owner_member_id": "M003",
            "status": "In progress", "review_status": "Pending",
        },
        source="tracker",
    )

    client = client_for("lead@nyu.edu")
    response = client.get("/tracker/?team=T001")

    assert response.status_code == 200
    assert b"Scoped project" in response.content
    assert b"Other team project" not in response.content
    assert b"Diabetes" in response.content
    assert b"IoC" not in response.content

    update = client.post(
        "/tracker/",
        {
            "action": "update", "table_name": "Milestones", "record_id": "MS001",
            "status": "Completed", "next_action": "Tampered",
        },
    )
    review = client.post(
        "/tracker/",
        {
            "action": "review", "review-record_type": "Milestone",
            "review-record_id": "MS001", "review-review_status": "Approved",
            "review-review_note": "Tampered",
        },
    )

    assert update.status_code == 403
    assert review.status_code == 403
    milestone = SheetRecord.objects.get(table_name="Milestones", record_id="MS001")
    assert milestone.payload["status"] == "In progress"
    assert milestone.payload["review_status"] == "Pending"


def test_tracker_only_shows_projects_assigned_to_members_team_or_member():
    add_record(
        "Members",
        "M002",
        {
            "member_id": "M002",
            "email": "member@nyu.edu",
            "display_name": "Assigned Member",
            "active": "TRUE",
        },
    )
    add_record(
        "Members",
        "M003",
        {
            "member_id": "M003",
            "email": "other@nyu.edu",
            "display_name": "Other Member",
            "active": "TRUE",
        },
    )
    for team_id, team_name in (("T001", "Core Lab"), ("T002", "Private Team")):
        add_record(
            "Teams",
            team_id,
            {"team_id": team_id, "team_name": team_name, "active": "TRUE"},
        )
    for member_id, team_id in (("M002", "T001"), ("M003", "T002")):
        add_record(
            "Member_Teams",
            f"MT-{member_id}",
            {
                "member_team_id": f"MT-{member_id}",
                "member_id": member_id,
                "team_id": team_id,
                "active": "TRUE",
            },
        )
    add_record(
        "App_Roles",
        "AR002",
        {
            "app_role_id": "AR002",
            "member_id": "M002",
            "app_id": "project_tracker",
            "app_role": "member",
            "scope_team_id": "",
            "active": "TRUE",
        },
    )
    projects = (
        (
            "P-DIRECT",
            "Directly assigned project",
            "T002",
            "M002",
        ),
        (
            "P-TEAM",
            "Team assigned project",
            "T001",
            "M003",
        ),
        (
            "P-HIDDEN",
            "Hidden private project",
            "T002",
            "M003",
        ),
    )
    for project_id, title, team_ids, member_ids in projects:
        add_record(
            "Projects",
            project_id,
            {
                "project_id": project_id,
                "project": title,
                "owner_member_id": "M003",
                "assigned_team_ids": team_ids,
                "assigned_member_ids": member_ids,
            },
            source="tracker",
        )
        add_record(
            "Milestones",
            f"MS-{project_id}",
            {
                "milestone_id": f"MS-{project_id}",
                "project_id": project_id,
                "milestone": f"{title} Gantt task",
                "owner_member_id": "M003",
                "start_date": "2026-09-01",
                "due_date": "2026-09-05",
                "status": "In progress",
                "review_status": "Approved" if project_id != "P-HIDDEN" else "Pending",
            },
            source="tracker",
        )

    client = client_for("member@nyu.edu")
    session = client.session
    session["gantt_import_preview"] = {
        "token": "hidden-token",
        "actor": "member@nyu.edu",
        "project_id": "P-HIDDEN",
        "project": "Hidden private project",
        "sheet_name": "hidden.xlsx",
        "rows": [
            {
                "project_id": "P-HIDDEN",
                "milestone": "Hidden preview task",
                "start_date": "2026-09-01",
                "due_date": "2026-09-05",
            }
        ],
        "warnings": [],
        "errors": [],
    }
    session.save()

    response = client.get("/tracker/?gantt_project=P-HIDDEN")

    assert response.status_code == 200
    assert b"Directly assigned project" in response.content
    assert b"Team assigned project" in response.content
    assert b"Hidden private project" not in response.content
    assert b"Hidden private project Gantt task" not in response.content
    assert b"Hidden preview task" not in response.content
    assert response.context["selected_gantt_project"]["project_id"] == "P-DIRECT"
    assert "gantt_import_preview" not in client.session

    update = client.post(
        "/tracker/",
        {
            "action": "update",
            "table_name": "Milestones",
            "record_id": "MS-P-HIDDEN",
            "status": "Completed",
            "next_action": "Tampered",
        },
    )
    assert update.status_code == 403
    assert SheetRecord.objects.get(
        source="tracker",
        table_name="Milestones",
        record_id="MS-P-HIDDEN",
    ).payload["status"] == "In progress"

    portal = client.get("/portal/")
    action_counts = {
        item["label"]: item["count"]
        for item in portal.context["action_panel"]["items"]
    }
    assert action_counts["Blocked"] == 0
    assert action_counts["Pending approval"] == 0


def test_portal_admin_can_view_all_project_gantt_charts():
    seed_pi()
    add_record(
        "Projects",
        "P-PRIVATE",
        {
            "project_id": "P-PRIVATE",
            "project": "Administratively visible project",
            "owner_member_id": "M099",
            "assigned_team_ids": "T099",
            "assigned_member_ids": "M099",
        },
        source="tracker",
    )
    add_record(
        "Milestones",
        "MS-PRIVATE",
        {
            "milestone_id": "MS-PRIVATE",
            "project_id": "P-PRIVATE",
            "milestone": "Administratively visible Gantt task",
            "owner_member_id": "M099",
            "start_date": "2026-09-01",
            "due_date": "2026-09-05",
        },
        source="tracker",
    )

    response = signed_in_client().get("/tracker/?gantt_project=P-PRIVATE")

    assert response.status_code == 200
    assert b"Administratively visible project" in response.content
    assert b"Administratively visible Gantt task" in response.content


@patch("labapps.views.upsert_record")
def test_project_assignments_save_multiple_teams_and_members(mock_upsert):
    seed_pi()
    add_record(
        "Members",
        "M002",
        {
            "member_id": "M002",
            "email": "member@nyu.edu",
            "display_name": "Research Member",
            "active": "TRUE",
        },
    )
    for team_id, team_name in (("T001", "Core Lab"), ("T002", "Diabetes")):
        add_record(
            "Teams",
            team_id,
            {"team_id": team_id, "team_name": team_name, "active": "TRUE"},
        )
    add_record(
        "Projects",
        "P001",
        {
            "project_id": "P001",
            "project": "Assigned project",
            "owner_member_id": "M001",
        },
        source="tracker",
    )

    response = signed_in_client().post(
        "/tracker/",
        {
            "action": "project_assignment",
            "project_id": "P001",
            "assignment-P001-assigned_team_ids": ["T001", "T002"],
            "assignment-P001-assigned_member_ids": ["M002"],
        },
    )

    assert response.status_code == 302
    assert response["Location"].endswith("/tracker/#projects")
    payload = mock_upsert.call_args.args[1]
    assert payload["assigned_team_ids"] == "T001;T002"
    assert payload["assigned_member_ids"] == "M002;M001"
    assert mock_upsert.call_args.kwargs["action"] == "update_project_assignments"


def test_project_assignment_choices_render_as_collapsible_dropdowns():
    seed_pi()
    add_record(
        "Teams",
        "T001",
        {"team_id": "T001", "team_name": "Core Lab", "active": "TRUE"},
    )
    add_record(
        "Projects",
        "P001",
        {
            "project_id": "P001",
            "project": "Assigned project",
            "owner_member_id": "M001",
            "assigned_team_ids": "T001",
            "assigned_member_ids": "M001",
        },
        source="tracker",
    )

    response = signed_in_client().get("/tracker/")

    assert response.status_code == 200
    assert response.content.count(b'data-assignment-dropdown="teams"') == 2
    assert response.content.count(b'data-assignment-dropdown="members"') == 2
    assert b'name="project-assigned_team_ids"' in response.content
    assert b'name="project-assigned_member_ids"' in response.content
    assert b'name="assignment-P001-assigned_team_ids"' in response.content
    assert b'name="assignment-P001-assigned_member_ids"' in response.content
    assert b"1 selected" in response.content
    assert b"labapps/assignment-dropdown.js" in response.content


@patch("labapps.views.upsert_record")
def test_scoped_assignment_edit_preserves_other_team_collaborators(mock_upsert):
    add_record(
        "Members",
        "M002",
        {
            "member_id": "M002",
            "email": "lead@nyu.edu",
            "display_name": "Diabetes Lead",
            "active": "TRUE",
        },
    )
    add_record(
        "Members",
        "M003",
        {
            "member_id": "M003",
            "email": "collaborator@nyu.edu",
            "display_name": "Core Collaborator",
            "active": "TRUE",
        },
    )
    for team_id, team_name in (("T001", "Core Lab"), ("T002", "Diabetes")):
        add_record(
            "Teams",
            team_id,
            {"team_id": team_id, "team_name": team_name, "active": "TRUE"},
        )
    for member_id, team_id in (("M002", "T002"), ("M003", "T001")):
        add_record(
            "Member_Teams",
            f"MT-{member_id}",
            {
                "member_team_id": f"MT-{member_id}",
                "member_id": member_id,
                "team_id": team_id,
                "active": "TRUE",
            },
        )
    add_record(
        "App_Roles",
        "AR002",
        {
            "app_role_id": "AR002",
            "member_id": "M002",
            "app_id": "project_tracker",
            "app_role": "lead",
            "scope_team_id": "T002",
            "active": "TRUE",
        },
    )
    add_record(
        "Projects",
        "P001",
        {
            "project_id": "P001",
            "project": "Cross-team project",
            "owner_member_id": "M002",
            "assigned_team_ids": "T001;T002",
            "assigned_member_ids": "M002;M003",
        },
        source="tracker",
    )

    response = client_for("lead@nyu.edu").post(
        "/tracker/?team=T002",
        {
            "action": "project_assignment",
            "project_id": "P001",
            "assignment-P001-assigned_team_ids": ["T002"],
            "assignment-P001-assigned_member_ids": ["M002"],
        },
    )

    assert response.status_code == 302
    payload = mock_upsert.call_args.args[1]
    assert set(payload["assigned_team_ids"].split(";")) == {"T001", "T002"}
    assert set(payload["assigned_member_ids"].split(";")) == {"M002", "M003"}


def test_team_filter_prefers_explicit_project_assignments_over_owner_team():
    add_record(
        "Members",
        "M002",
        {"member_id": "M002", "email": "lead@nyu.edu", "active": "TRUE"},
    )
    add_record(
        "Members",
        "M003",
        {"member_id": "M003", "email": "other@nyu.edu", "active": "TRUE"},
    )
    for team_id, team_name in (("T001", "Core Lab"), ("T002", "Diabetes")):
        add_record(
            "Teams",
            team_id,
            {"team_id": team_id, "team_name": team_name, "active": "TRUE"},
        )
    for member_id, team_id in (("M002", "T002"), ("M003", "T001")):
        add_record(
            "Member_Teams",
            f"MT-{member_id}",
            {
                "member_team_id": f"MT-{member_id}",
                "member_id": member_id,
                "team_id": team_id,
                "active": "TRUE",
            },
        )
    add_record(
        "App_Roles",
        "AR002",
        {
            "app_role_id": "AR002",
            "member_id": "M002",
            "app_id": "project_tracker",
            "app_role": "lead",
            "scope_team_id": "T002",
            "active": "TRUE",
        },
    )
    add_record(
        "Projects",
        "P-CROSS",
        {
            "project_id": "P-CROSS",
            "project": "Explicit Diabetes project",
            "owner_member_id": "M003",
            "assigned_team_ids": "T002",
            "assigned_member_ids": "M002;M003",
        },
        source="tracker",
    )
    add_record(
        "Projects",
        "P-OTHER",
        {
            "project_id": "P-OTHER",
            "project": "Explicit Core project",
            "owner_member_id": "M002",
            "assigned_team_ids": "T001",
        },
        source="tracker",
    )
    add_record(
        "Milestones",
        "MS-OTHER",
        {
            "milestone_id": "MS-OTHER",
            "project_id": "P-OTHER",
            "milestone": "Hidden Core milestone",
            "owner_member_id": "M002",
            "status": "In progress",
            "review_status": "Pending",
        },
        source="tracker",
    )

    response = client_for("lead@nyu.edu").get("/tracker/?team=T002")

    assert response.status_code == 200
    assert b"Explicit Diabetes project" in response.content
    assert b"Explicit Core project" not in response.content
    assert b"Hidden Core milestone" not in response.content
    assert b"other@nyu.edu" not in response.content


def test_project_assignment_is_reassigned_before_member_deletion(
    monkeypatch,
    settings,
):
    settings.ENABLE_SHEET_WRITES = True
    settings.SHEET_WRITE_ALLOWED_EMAILS = {"*"}
    add_record(
        "Members",
        "M001",
        {
            "member_id": "M001",
            "email": "owner@nyu.edu",
            "global_role": "admin",
            "active": "TRUE",
        },
    )
    add_record(
        "Members",
        "M020",
        {
            "member_id": "M020",
            "email": "former@nyu.edu",
            "global_role": "member",
            "active": "FALSE",
        },
    )
    add_record(
        "Projects",
        "P020",
        {
            "project_id": "P020",
            "owner_member_id": "M001",
            "assigned_member_ids": "M020;M001",
        },
        source="tracker",
    )
    monkeypatch.setattr("labapps.services.members.replace_table", fake_registry_replace)
    monkeypatch.setattr(
        "labapps.services.members.append_registry_audit", lambda **kwargs: None
    )

    reference = next(
        row
        for row in member_reference_summary("M020")
        if row["table_name"] == "Project assignments"
    )
    assert reference["count"] == 1

    delete_member_record(
        "M020",
        actor="admin@nyu.edu",
        confirm_email="former@nyu.edu",
        reassign_to_member_id="M001",
    )

    project = SheetRecord.objects.get(
        source="tracker", table_name="Projects", record_id="P020"
    )
    assert project.payload["assigned_member_ids"] == "M001"


def test_tracker_member_with_two_scoped_roles_can_switch_between_both_teams():
    add_record(
        "Members",
        "M002",
        {
            "member_id": "M002",
            "email": "multiteam@nyu.edu",
            "display_name": "Multi Team",
            "active": "TRUE",
        },
    )
    for team_id, team_name in (("T001", "IoC"), ("T002", "Diabetes")):
        add_record(
            "Teams",
            team_id,
            {"team_id": team_id, "team_name": team_name, "active": "TRUE"},
        )
        add_record(
            "Member_Teams",
            f"MT-{team_id}",
            {
                "member_team_id": f"MT-{team_id}",
                "member_id": "M002",
                "team_id": team_id,
                "active": "TRUE",
            },
        )
        add_record(
            "App_Roles",
            f"AR-{team_id}",
            {
                "member_id": "M002",
                "app_id": "project_tracker",
                "app_role": "lead",
                "scope_team_id": team_id,
                "active": "TRUE",
            },
        )
        add_record(
            "Projects",
            f"P-{team_id}",
            {
                "project_id": f"P-{team_id}",
                "project": f"{team_name} project",
                "owner_member_id": "M002",
            },
            source="tracker",
        )

    client = client_for("multiteam@nyu.edu")
    ioc = client.get("/tracker/?team=T001")
    diabetes = client.get("/tracker/?team=T002")

    assert ioc.status_code == diabetes.status_code == 200
    assert b"IoC project" in ioc.content
    assert b"Diabetes project" in diabetes.content


def test_tracker_viewer_scope_stays_read_only_when_another_scope_is_lead():
    add_record(
        "Members",
        "M002",
        {
            "member_id": "M002",
            "email": "mixed@nyu.edu",
            "display_name": "Mixed Role",
            "active": "TRUE",
        },
    )
    add_record(
        "Members",
        "M003",
        {
            "member_id": "M003",
            "email": "owner@nyu.edu",
            "display_name": "Owner",
            "active": "TRUE",
        },
    )
    for team_id, team_name, role in (
        ("T001", "IoC", "viewer"),
        ("T002", "Diabetes", "lead"),
    ):
        add_record(
            "Teams",
            team_id,
            {"team_id": team_id, "team_name": team_name, "active": "TRUE"},
        )
        add_record(
            "App_Roles",
            f"AR-{team_id}",
            {
                "member_id": "M002",
                "app_id": "project_tracker",
                "app_role": role,
                "scope_team_id": team_id,
                "active": "TRUE",
            },
        )
    add_record(
        "Member_Teams",
        "MT-OWNER",
        {
            "member_team_id": "MT-OWNER",
            "member_id": "M003",
            "team_id": "T001",
            "active": "TRUE",
        },
    )
    add_record(
        "Milestones",
        "MS-VIEW",
        {
            "milestone_id": "MS-VIEW",
            "milestone": "Viewer-only milestone",
            "owner_member_id": "M003",
            "status": "In progress",
            "review_status": "Pending",
        },
        source="tracker",
    )

    response = client_for("mixed@nyu.edu").post(
        "/tracker/?team=T001",
        {
            "action": "update",
            "table_name": "Milestones",
            "record_id": "MS-VIEW",
            "status": "Completed",
            "next_action": "Tampered",
        },
    )

    assert response.status_code == 403
    assert SheetRecord.objects.get(
        table_name="Milestones",
        record_id="MS-VIEW",
    ).payload["status"] == "In progress"


def test_tracker_rejects_unsafe_experiment_data_url():
    seed_pi()
    add_record(
        "Experiments",
        "EXP-URL",
        {
            "experiment_id": "EXP-URL",
            "experiment_title": "URL validation",
            "member_id": "M001",
            "status": "In progress",
            "review_status": "Pending",
            "experiment_data_link": "https://example.com/data",
        },
        source="tracker",
    )

    response = signed_in_client().post(
        "/tracker/",
        {
            "action": "update",
            "table_name": "Experiments",
            "record_id": "EXP-URL",
            "status": "In progress",
            "next_action": "Keep safe",
            "experiment_data_link": "javascript:alert(1)",
        },
    )

    assert response.status_code == 200
    assert SheetRecord.objects.get(
        table_name="Experiments",
        record_id="EXP-URL",
    ).payload["experiment_data_link"] == "https://example.com/data"


@patch("labapps.views.store_knowledge_file", return_value=("knowledge/N1/file.pdf", "abc123"))
def test_private_knowledge_upload_creates_record(mock_store):
    seed_pi()
    client = signed_in_client()
    response = client.post(
        "/knowledge/upload/",
        {
            "record_type": "protocol", "title": "New protocol", "team": "IoC",
            "owner": "Ken", "category": "Assay", "notes": "verified",
            "files": SimpleUploadedFile("protocol.pdf", b"PDF", content_type="application/pdf"),
        },
    )
    assert response.status_code == 302
    record = KnowledgeRecord.objects.get(title="New protocol")
    assert record.object_name == "knowledge/N1/file.pdf"
    assert record.metadata["sha256"] == "abc123"
    assert record.metadata["parse_status"] == "failed"
    mock_store.assert_called_once()


@patch(
    "labapps.views.store_knowledge_file",
    return_value=("knowledge/P1/mef.docx", "docx-sha"),
)
def test_protocol_upload_extracts_and_displays_structured_content(mock_store):
    seed_pi()
    client = signed_in_client()
    content = protocol_docx_bytes()

    response = client.post(
        "/knowledge/upload/",
        {
            "record_type": "protocol",
            "title": "MEF preparation protocol",
            "team": "Common",
            "owner": "Ken",
            "category": "Cell culture",
            "notes": "Uploaded from the lab template",
            "files": SimpleUploadedFile(
                "MEF_protocol.docx",
                content,
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        },
    )

    assert response.status_code == 302
    record = KnowledgeRecord.objects.get(title="MEF preparation protocol")
    assert record.metadata["parse_status"] == "parsed"
    assert record.metadata["section_count"] == 3
    assert record.metadata["procedure"] == [
        "Collect each embryo separately.",
        "Plate cells in complete medium.",
    ]
    assert mock_store.call_args.args[2] == content

    detail = client.get(f"/knowledge/?protocol={record.record_id}")
    assert b"Prepare primary MEFs from individual embryos." in detail.content
    assert b"DMEM" in detail.content
    assert b"Collect each embryo separately." in detail.content
    assert f"/knowledge/{record.record_id}/reprocess/".encode() in detail.content


def test_upload_deletes_new_object_when_database_persistence_fails():
    seed_pi()
    client = signed_in_client()
    with (
        patch(
            "labapps.views.store_knowledge_file",
            return_value=("knowledge/P-ORPHAN/protocol.docx", "stored-sha"),
        ),
        patch(
            "labapps.views.KnowledgeRecord.objects.create",
            side_effect=RuntimeError("database unavailable"),
        ),
        patch("labapps.views.delete_knowledge_file") as mock_delete,
    ):
        response = client.post(
            "/knowledge/upload/",
            {
                "record_type": "protocol",
                "title": "Atomic protocol",
                "team": "Common",
                "owner": "Ken",
                "category": "Cell culture",
                "notes": "",
                "files": SimpleUploadedFile(
                    "protocol.docx",
                    protocol_docx_bytes(),
                    content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
            },
        )

    assert response.status_code == 200
    assert b"database unavailable" in response.content
    mock_delete.assert_called_once_with(
        "knowledge/P-ORPHAN/protocol.docx"
    )


@patch("labapps.views.read_knowledge_file", return_value=protocol_docx_bytes())
def test_existing_protocol_can_be_reprocessed_without_losing_metadata(mock_read):
    seed_pi()
    record = KnowledgeRecord.objects.create(
        record_id="P-MEF",
        record_type="protocol",
        title="MEF preparation protocol",
        team="Common",
        owner="Ken",
        status="active",
        object_name="knowledge/P-MEF/mef.docx",
        original_filename="MEF_protocol.docx",
        metadata={
            "notes": "Keep this note",
            "sha256": hashlib.sha256(mock_read.return_value).hexdigest(),
            "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        },
    )
    client = signed_in_client()

    response = client.post(f"/knowledge/{record.record_id}/reprocess/")

    assert response.status_code == 302
    record.refresh_from_db()
    assert record.metadata["notes"] == "Keep this note"
    assert record.metadata["sha256"] == hashlib.sha256(
        mock_read.return_value
    ).hexdigest()
    assert record.metadata["parse_status"] == "parsed"
    assert record.metadata["section_count"] == 3
    mock_read.assert_called_once_with("knowledge/P-MEF/mef.docx")
    assert LabAppAudit.objects.filter(
        target="P-MEF",
        action="reprocess_content",
    ).exists()


@patch("labapps.views.read_knowledge_file", return_value=b"changed")
def test_reprocess_stops_when_original_checksum_changed(mock_read):
    seed_pi()
    record = KnowledgeRecord.objects.create(
        record_id="P-CHECKSUM",
        record_type="protocol",
        title="Checksum protocol",
        object_name="knowledge/P-CHECKSUM/protocol.docx",
        original_filename="protocol.docx",
        metadata={"sha256": hashlib.sha256(b"original").hexdigest()},
    )

    response = signed_in_client().post(
        f"/knowledge/{record.record_id}/reprocess/"
    )

    assert response.status_code == 302
    record.refresh_from_db()
    assert "sections" not in record.metadata
    assert LabAppAudit.objects.filter(
        target="P-CHECKSUM",
        action="reprocess_checksum_mismatch",
    ).exists()
    mock_read.assert_called_once()


@patch("labapps.views.read_knowledge_file", return_value=b"broken-docx")
def test_failed_reprocess_preserves_old_content_and_shows_stale_warning(mock_read):
    seed_pi()
    record = KnowledgeRecord.objects.create(
        record_id="P-STALE",
        record_type="protocol",
        title="Previously parsed protocol",
        object_name="knowledge/P-STALE/protocol.docx",
        original_filename="protocol.docx",
        metadata={
            "sha256": hashlib.sha256(b"broken-docx").hexdigest(),
            "parse_status": "parsed",
            "section_count": 1,
            "sections": [
                {
                    "heading": "Procedure",
                    "blocks": [
                        {"kind": "paragraph", "text": "Previously extracted step"}
                    ],
                }
            ],
        },
    )
    client = signed_in_client()

    response = client.post(f"/knowledge/{record.record_id}/reprocess/")

    assert response.status_code == 302
    record.refresh_from_db()
    assert record.metadata["parse_status"] == "parsed"
    assert record.metadata["last_reprocess_status"] == "failed"
    assert record.metadata["sections"][0]["heading"] == "Procedure"

    detail = client.get(f"/knowledge/?protocol={record.record_id}")
    assert b"Previously extracted step" in detail.content
    assert b"latest reprocessing attempt failed" in detail.content
    mock_read.assert_called_once()


@patch("labapps.views.read_knowledge_file")
def test_read_only_member_cannot_reprocess_protocol(mock_read):
    add_record(
        "Members",
        "M002",
        {
            "member_id": "M002",
            "email": "member@nyu.edu",
            "display_name": "Lab member",
            "global_role": "member",
            "active": "TRUE",
        },
    )
    add_record(
        "App_Roles",
        "AR-notebooks-reader",
        {
            "member_id": "M002",
            "app_id": "notebooks_protocols",
            "app_role": "viewer",
            "scope_team_id": "",
            "active": "TRUE",
        },
    )
    KnowledgeRecord.objects.create(
        record_id="P-LOCKED",
        record_type="protocol",
        title="Locked protocol",
        object_name="knowledge/P-LOCKED/protocol.docx",
        original_filename="protocol.docx",
    )

    response = client_for("member@nyu.edu").post(
        "/knowledge/P-LOCKED/reprocess/"
    )

    assert response.status_code == 403
    mock_read.assert_not_called()
